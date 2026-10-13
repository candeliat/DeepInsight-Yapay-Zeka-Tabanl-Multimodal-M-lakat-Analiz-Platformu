"""
CV dosyasından metin çıkarma. Sıra: (1) PDF metin katmanı (PyMuPDF), (2) metin
yoksa ya da dosya görselse OCR (EasyOCR, Türkçe+İngilizce), (3) DOCX.

Güvenlik: dosya türü istemcinin bildirdiği Content-Type'a değil, dosyanın ilk
baytlarına (magic bytes) bakılarak belirlenir. Çıkarılan metin ham VERİDİR —
LLM'e verilirken talimat olarak yorumlanmamalıdır (bkz. ileride cv_service).

Tüm fonksiyonlar senkron ve CPU-bound'dur; çağıran taraf `asyncio.to_thread`
ile çağırmalıdır.
"""
import io
import logging
import re
import threading
import zipfile
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

MAX_CV_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 30_000
MAX_OCR_PAGES = 4
# PDF'in metin katmanı bu kadar karakterden azsa taranmış sayılır ve OCR'a düşülür.
MIN_TEXT_LAYER_CHARS = 200
# Bu boyuttan küçük görseller (madde işareti, küçük ikon) fotoğraf/grafik sayılmaz (pt).
MIN_GRAPHIC_SIDE = 40
# Bundan küçük yazı okunamaz; ATS'lere anahtar kelime gizlemek için kullanılır (pt).
MIN_VISIBLE_FONT = 3
_WHITE = 0xFFFFFF

_ocr_reader = None
_ocr_lock = threading.Lock()


class UnsupportedCVFile(ValueError):
    pass


def detect_kind(data: bytes) -> Tuple[str, str]:
    """(kind, mime) döner; kind: 'pdf' | 'image' | 'docx'. Desteklenmeyen dosyada UnsupportedCVFile."""
    if data[:5] == b"%PDF-":
        return "pdf", "application/pdf"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image", "image/png"
    if data[:3] == b"\xff\xd8\xff":
        return "image", "image/jpeg"
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except zipfile.BadZipFile:
            pass
    raise UnsupportedCVFile("Desteklenmeyen dosya türü. PDF, DOCX, PNG veya JPEG yükleyin.")


def _get_ocr_reader():
    global _ocr_reader
    if _ocr_reader is None:
        with _ocr_lock:
            if _ocr_reader is None:
                import os
                import easyocr
                logger.info("EasyOCR (tr, en) yükleniyor...")
                # Model dosyalarının (~100 MB) varsayılan olarak kullanıcı profiline
                # inmemesi için dizin ortam değişkeniyle değiştirilebilir.
                _ocr_reader = easyocr.Reader(
                    ["tr", "en"], gpu=False, verbose=False,
                    model_storage_directory=os.getenv("EASYOCR_MODEL_DIR") or None,
                )
    return _ocr_reader


def _ocr_image_bytes(image_bytes: bytes) -> str:
    # detail=0 -> yalnızca metin; paragraph=True satırları birleştirir.
    lines = _get_ocr_reader().readtext(image_bytes, detail=0, paragraph=True)
    return "\n".join(lines)


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def _estimate_columns(page) -> int:
    """Sayfadaki metin bloklarının x konumlarına göre kaba sütun tahmini (1 ya da 2+).
    ATS ayrıştırıcıları çok sütunlu düzeni sıklıkla karıştırdığı için skorda kullanılır."""
    width = page.rect.width
    blocks = [b for b in page.get_text("blocks") if b[4].strip()]
    if len(blocks) < 6:
        return 1
    left = sum(1 for b in blocks if b[2] < width * 0.55)
    right = sum(1 for b in blocks if b[0] > width * 0.45)
    both_sides = left >= 3 and right >= 3 and not any(b[0] < width * 0.3 and b[2] > width * 0.7 for b in blocks[:3])
    return 2 if both_sides else 1


def _overlaps(a, b) -> bool:
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def _pdf_page_layout(page) -> Tuple[int, List[str]]:
    """(grafik sayısı, gizli metin parçaları) döner.

    Grafik: kenarları MIN_GRAPHIC_SIDE'dan büyük görseller (fotoğraf, logo, grafik).
    Gizli metin: beyaz, tamamen saydam ya da MIN_VISIBLE_FONT'tan küçük yazı — ama
    renkli bir dolgu ya da görselin üzerindeyse değil (koyu kenar çubuğundaki beyaz
    başlık, taranmış sayfanın üstündeki görünmez OCR katmanı meşrudur).
    """
    images = [i["bbox"] for i in page.get_image_info()]
    graphics = sum(1 for b in images if b[2] - b[0] >= MIN_GRAPHIC_SIDE and b[3] - b[1] >= MIN_GRAPHIC_SIDE)
    backgrounds = list(images)
    for d in page.get_drawings():
        fill = d.get("fill")
        if fill and tuple(fill) != (1.0, 1.0, 1.0):
            backgrounds.append(tuple(d["rect"]))
    hidden: List[str] = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                if not span["text"].strip():
                    continue
                invisible = span["color"] == _WHITE or span.get("alpha", 255) == 0 or span["size"] < MIN_VISIBLE_FONT
                if invisible and not any(_overlaps(span["bbox"], bg) for bg in backgrounds):
                    hidden.append(span["text"])
    return graphics, hidden


def _extract_pdf(data: bytes) -> Tuple[str, str, Dict[str, Any]]:
    import pymupdf as fitz

    with fitz.open(stream=data, filetype="pdf") as doc:
        page_count = doc.page_count
        columns = max((_estimate_columns(p) for p in doc), default=1)
        layout = [_pdf_page_layout(p) for p in doc]
        # Gizli metin profile/ilan eşleşmesine girmesin: ATS'yi kandırmaya yönelik anahtar
        # kelimeler adayın becerisi sayılmaz (yalnızca skorda "gizli metin" olarak raporlanır).
        pages = []
        for page, (_, hidden) in zip(doc, layout):
            page_text = page.get_text("text")
            for fragment in hidden:
                page_text = page_text.replace(fragment, "", 1)
            pages.append(page_text)
        text = "\n".join(pages)
        stats: Dict[str, Any] = {
            "page_count": page_count, "columns": columns, "has_text_layer": len(text.strip()) >= MIN_TEXT_LAYER_CHARS,
            "graphic_count": sum(g for g, _ in layout),
            "hidden_text_chars": sum(len(f.strip()) for _, hidden in layout for f in hidden),
            # PDF'te üst/alt bilgi ayrı bir alan değildir; ayrıştırıcı onu sayfa metni olarak okur.
            "header_footer_contact": False,
        }

        if stats["has_text_layer"]:
            return _clean(text), "pdf_text", stats

        # Taranmış PDF: sayfa görüntüsü CV'nin kendisidir, fotoğraf/grafik sayılmaz
        # (okunamazlığı zaten "metin katmanı" kontrolü cezalandırıyor).
        stats["graphic_count"] = 0
        # Sayfaları görüntüye çevirip OCR.
        parts = []
        for page in list(doc)[:MAX_OCR_PAGES]:
            pix = page.get_pixmap(dpi=200)
            parts.append(_ocr_image_bytes(pix.tobytes("png")))
        return _clean("\n".join(parts)), "ocr", stats


def _docx_header_footer_text(document) -> str:
    parts = []
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part.is_linked_to_previous:
                continue
            parts.extend(p.text for p in part.paragraphs)
            for table in part.tables:
                parts.extend(c.text for row in table.rows for c in row.cells)
    return "\n".join(t for t in parts if t.strip())


def _docx_run_is_hidden(run) -> bool:
    from docx.shared import Pt

    font = run.font
    color = font.color.rgb if font.color is not None and font.color.type is not None else None
    return bool(font.hidden or str(color) == "FFFFFF" or (font.size is not None and font.size < Pt(MIN_VISIBLE_FONT)))


def _docx_visible_paragraphs(document) -> Tuple[List[str], int]:
    """(görünür paragraf metinleri, gizli karakter sayısı). Gizli koşular metne alınmaz:
    ATS'yi kandırmaya yönelik anahtar kelimeler profile/ilan eşleşmesine girmez."""
    visible, hidden = [], 0
    for paragraph in document.paragraphs:
        parts = []
        for run in paragraph.runs:
            if run.text.strip() and _docx_run_is_hidden(run):
                hidden += len(run.text.strip())
            else:
                parts.append(run.text)
        # Koşu dışı metin (köprü vb.) kaybolmasın diye gizli koşu yoksa paragraf metni aynen alınır.
        line = "".join(parts) if len(parts) != len(paragraph.runs) else paragraph.text
        if line.strip():
            visible.append(line)
    return visible, hidden


def _extract_docx(data: bytes) -> Tuple[str, str, Dict[str, Any]]:
    import docx
    from docx.oxml.ns import qn

    from app.services.ats_service import find_contact

    document = docx.Document(io.BytesIO(data))
    paragraphs, hidden_chars = _docx_visible_paragraphs(document)
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    text = "\n".join(paragraphs)
    # Workday gibi ayrıştırıcılar üst/alt bilgiyi okumaz; oradaki iletişim bilgisi kaybolur.
    # Metnin kendisi saklanmaz (kişisel veri), yalnızca orada iletişim bilgisi olup olmadığı.
    header_contact = find_contact(_docx_header_footer_text(document))
    body = document.element.body
    stats = {
        "page_count": None, "columns": 1, "has_text_layer": True, "table_count": len(document.tables),
        "graphic_count": sum(1 for _ in body.iter(qn("w:drawing"))) + sum(1 for _ in body.iter(qn("w:pict"))),
        "hidden_text_chars": hidden_chars,
        "header_footer_contact": bool(header_contact["email"] or header_contact["phone"]),
    }
    return _clean(text), "docx", stats


def extract_cv_text(data: bytes) -> Tuple[str, str, str, Dict[str, Any]]:
    """
    (metin, yöntem, mime, istatistikler) döner. Dosya boyutu/türü geçersizse
    UnsupportedCVFile fırlatır. Metin boşsa (ör. OCR hiçbir şey okuyamadı)
    çağıran taraf bunu 'failed' olarak işaretlemelidir.
    """
    if len(data) > MAX_CV_BYTES:
        raise UnsupportedCVFile(f"Dosya çok büyük (en fazla {MAX_CV_BYTES // (1024 * 1024)} MB).")
    kind, mime = detect_kind(data)

    if kind == "pdf":
        text, method, stats = _extract_pdf(data)
    elif kind == "docx":
        text, method, stats = _extract_docx(data)
    else:
        text = _clean(_ocr_image_bytes(data))
        method, stats = "ocr", {"page_count": 1, "columns": 1, "has_text_layer": False}

    stats["char_count"] = len(text)
    stats["file_size"] = len(data)
    return text, method, mime, stats
