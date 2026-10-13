"""
ATS ayrıştırma riskleri (docs/research/MULAKAT_VE_ATS_ARASTIRMASI.md §4.2): fotoğraf/grafik,
2,5 MB sınırı, üst/alt bilgideki iletişim bilgisi, gizli metin, açılımsız kısaltmalar.
İstatistikler gerçek PDF/DOCX dosyaları üretilerek çıkarılır; OCR ve ağ yok.
"""
import io

import docx
import pymupdf
from docx.shared import Pt, RGBColor

from app.services import ats_service, cv_text_service

BODY = (
    "Ayse Yilmaz\nDENEYIM\nBackend Developer (2021 - 2023)\nEGITIM\nITU Bilgisayar Muhendisligi (2016 - 2020)\n"
    "BECERILER\nPython, FastAPI, PostgreSQL, Docker\n" + "Proje aciklamasi ve sorumluluklar. " * 20
)
STATS = {"page_count": 1, "columns": 1, "char_count": 2400, "has_text_layer": True}


def _by_id(result):
    return {c["id"]: c for c in result["breakdown"]}


def _png(size=120):
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, size, size), 0)
    pix.set_rect(pix.irect, (200, 30, 30))
    return pix.tobytes("png")


def _pdf(build):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), BODY, fontsize=10)
    build(page)
    return doc.tobytes()


def _docx(build):
    document = docx.Document()
    for line in BODY.splitlines():
        document.add_paragraph(line)
    build(document)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# --- metin çıkarma istatistikleri ---

def test_pdf_photo_is_counted_but_small_icon_is_not():
    def build(page):
        page.insert_image(pymupdf.Rect(450, 40, 550, 140), stream=_png())  # fotoğraf
        page.insert_image(pymupdf.Rect(40, 40, 50, 50), stream=_png(16))   # madde işareti boyutunda ikon

    _, _, _, stats = cv_text_service.extract_cv_text(_pdf(build))
    assert stats["graphic_count"] == 1
    assert stats["file_size"] > 0


def test_pdf_white_text_on_white_is_hidden_but_white_on_dark_sidebar_is_not():
    def white_keywords(page):
        page.insert_text((50, 820), "Kubernetes Terraform Kafka Spark Airflow", fontsize=10, color=(1, 1, 1))

    def dark_sidebar(page):
        page.draw_rect(pymupdf.Rect(0, 800, 600, 840), fill=(0.1, 0.2, 0.4))
        page.insert_text((50, 820), "Kubernetes Terraform Kafka Spark Airflow", fontsize=10, color=(1, 1, 1))

    assert cv_text_service.extract_cv_text(_pdf(white_keywords))[3]["hidden_text_chars"] >= 30
    assert cv_text_service.extract_cv_text(_pdf(dark_sidebar))[3]["hidden_text_chars"] == 0


def test_pdf_tiny_font_text_is_hidden():
    def tiny(page):
        page.insert_text((50, 825), "react angular vue svelte nextjs nuxt", fontsize=1)

    assert cv_text_service.extract_cv_text(_pdf(tiny))[3]["hidden_text_chars"] >= 20


def test_docx_contact_only_in_header_is_detected_and_not_stored():
    def header(document):
        document.sections[0].header.paragraphs[0].text = "ayse@example.com | +90 555 123 45 67"

    text, method, _, stats = cv_text_service.extract_cv_text(_docx(header))
    assert method == "docx" and stats["header_footer_contact"] is True
    assert "ayse@example.com" not in text and "ayse@example.com" not in str(stats)

    checks = _by_id(ats_service.score_cv(text, stats, method))
    assert checks["header_footer"]["status"] == "fail" and "üst bilgi" in checks["header_footer"]["tip"]
    assert checks["contact"]["status"] == "fail"


def test_docx_hidden_and_white_runs_and_pictures_are_counted():
    def build(document):
        hidden = document.add_paragraph().add_run("Kubernetes Terraform Kafka")
        hidden.font.hidden = True
        white = document.add_paragraph().add_run("Spark Airflow Snowflake")
        white.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tiny = document.add_paragraph().add_run("normal gorunen ama minik")
        tiny.font.size = Pt(2)
        document.add_picture(io.BytesIO(_png()))

    _, _, _, stats = cv_text_service.extract_cv_text(_docx(build))
    assert stats["hidden_text_chars"] >= 60
    assert stats["graphic_count"] == 1
    assert stats["header_footer_contact"] is False


# --- skor ---

def test_file_over_2_5_mb_fails_size_check_and_unknown_size_is_not_penalized():
    big = _by_id(ats_service.score_cv(BODY, {**STATS, "file_size": 3 * 1024 * 1024}, "pdf_text"))["file_size"]
    assert big["status"] == "fail" and "2,5 MB" in big["detail"]
    old = _by_id(ats_service.score_cv(BODY, STATS, "pdf_text"))["file_size"]
    assert old["status"] == "pass"


def test_graphics_reduce_formatting_score():
    check = _by_id(ats_service.score_cv(BODY, {**STATS, "graphic_count": 1}, "pdf_text"))["formatting"]
    assert check["status"] == "partial" and "1 fotoğraf/grafik" in check["detail"]


def test_hidden_text_fails_only_above_threshold():
    assert _by_id(ats_service.score_cv(BODY, {**STATS, "hidden_text_chars": 120}, "pdf_text"))["hidden_text"]["status"] == "fail"
    assert _by_id(ats_service.score_cv(BODY, {**STATS, "hidden_text_chars": 3}, "pdf_text"))["hidden_text"]["status"] == "pass"


def test_partial_contact_in_header_is_partial():
    text = "ayse@example.com\n" + BODY
    check = _by_id(ats_service.score_cv(text, {**STATS, "header_footer_contact": True}, "docx"))["header_footer"]
    assert check["status"] == "partial"


def test_acronym_without_expansion_is_suggested_with_expansion_first():
    text = BODY + "\nAWS, CI/CD ve ML projeleri. SQL ve HTML biliyorum."
    check = _by_id(ats_service.score_cv(text, STATS, "pdf_text"))["acronyms"]
    assert check["status"] == "partial"
    assert "AWS" in check["detail"] and "SQL" not in check["detail"]
    assert "Amazon Web Services (AWS)" in check["tip"]


def test_acronym_with_expansion_or_turkish_equivalent_passes():
    text = BODY + "\nAmazon Web Services (AWS) ve makine öğrenmesi (ML) deneyimi."
    assert ats_service.missing_acronym_expansions(text) == []


def test_acronym_inside_other_words_is_not_matched():
    assert ats_service.missing_acronym_expansions("HTML5, MAIL, EMLAK, BAWSE, ci/cd-like") == []


def test_legacy_upload_without_new_stats_is_not_penalized_and_says_unknown():
    checks = _by_id(ats_service.score_cv(BODY, STATS, "pdf_text"))
    for check_id in ("file_size", "header_footer", "hidden_text"):
        assert checks[check_id]["status"] == "pass" and "Bilinmiyor" in checks[check_id]["detail"]
    assert "bilinmiyor" in checks["formatting"]["detail"]


def test_hidden_keywords_are_removed_from_extracted_text():
    # ATS'yi kandırmaya yönelik metin profil/ilan eşleşmesine girmemeli; görünür metin korunur.
    def pdf_white(page):
        page.insert_text((50, 820), "Kubernetes Terraform Kafka", fontsize=10, color=(1, 1, 1))

    def docx_hidden(document):
        p = document.add_paragraph("Ek beceriler: Go ")
        run = p.add_run("Kubernetes Terraform Kafka")
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    pdf_text = cv_text_service.extract_cv_text(_pdf(pdf_white))[0]
    docx_text = cv_text_service.extract_cv_text(_docx(docx_hidden))[0]
    for text in (pdf_text, docx_text):
        assert "Kubernetes" not in text and "FastAPI" in text
    assert "Ek beceriler: Go" in docx_text


def test_contact_in_header_gives_one_suggestion_not_two():
    result = ats_service.score_cv(BODY, {**STATS, "graphic_count": 0, "header_footer_contact": True}, "docx")
    checks = _by_id(result)
    assert "üst/alt bilgide var" in checks["contact"]["detail"]
    assert checks["contact"]["tip"] == checks["header_footer"]["tip"]
    assert result["suggestions"].count(checks["contact"]["tip"]) == 1


def test_small_file_size_is_shown_in_kb():
    detail = _by_id(ats_service.score_cv(BODY, {**STATS, "file_size": 38_000}, "docx"))["file_size"]["detail"]
    assert detail.startswith("37 KB")
